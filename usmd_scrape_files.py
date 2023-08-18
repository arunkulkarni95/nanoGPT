import os
import requests
import PyPDF2
from docx import Document
import openpyxl
from urllib.parse import urljoin
from bs4 import BeautifulSoup

# Base URLs to scrape
base_urls = ["https://www.usmd.edu", "https://www.usmd.edu/usm"]

# Directory to save text files
output_directory = "scrape_results/usmd/files"
os.makedirs(output_directory, exist_ok=True)

# Function to scrape text from a PDF
def scrape_pdf(pdf_url, output_path):
    response = requests.get(pdf_url)
    with open(output_path, "wb") as pdf_file:
        pdf_file.write(response.content)
    pdf_text = ""
    with open(output_path, "rb") as pdf_file:
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        for page in pdf_reader.pages:
            pdf_text += page.extract_text()
    return pdf_text

# Function to scrape text from a Word document
def scrape_docx(docx_url, output_path):
    response = requests.get(docx_url)
    with open(output_path, "wb") as docx_file:
        docx_file.write(response.content)
    docx_text = ""
    doc = Document(output_path)
    for paragraph in doc.paragraphs:
        docx_text += paragraph.text + "\n"
    return docx_text

# Function to scrape text from an Excel spreadsheet
def scrape_xlsx(xlsx_url, output_path):
    response = requests.get(xlsx_url)
    with open(output_path, "wb") as xlsx_file:
        xlsx_file.write(response.content)
    xlsx_text = ""
    workbook = openpyxl.load_workbook(output_path)
    worksheet = workbook.active
    for row in worksheet.iter_rows(values_only=True):
        for cell_value in row:
            if cell_value:
                xlsx_text += str(cell_value) + "\n"
    return xlsx_text

# Loop through base URLs
for base_url in base_urls:
    response = requests.get(base_url)
    soup = BeautifulSoup(response.content, "html.parser")

    # Find and scrape links to PDFs, Word docs, and Excel sheets
    for link in soup.find_all("a"):
        link_url = link.get("href")
        absolute_link = urljoin(base_url, link_url)
        if link_url.lower().endswith((".pdf", ".docx", ".xlsx")):
            file_extension = link_url.split(".")[-1]
            file_name = link_url.split("/")[-1]
            output_path = os.path.join(output_directory, f"{file_name}.{file_extension}.txt")
            if file_extension == "pdf":
                scraped_text = scrape_pdf(absolute_link, output_path)
            elif file_extension == "docx":
                scraped_text = scrape_docx(absolute_link, output_path)
            elif file_extension == "xlsx":
                scraped_text = scrape_xlsx(absolute_link, output_path)
            with open(output_path, "w", encoding="utf-8") as txt_file:
                txt_file.write(scraped_text)
            print(f"Saved {file_name}.{file_extension}.txt")

print("Text extraction is complete.")
